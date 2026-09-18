import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Configuración de márgenes
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Estilo de Título Principal
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run("SISTEMA INTEGRAL ERP DE GESTIÓN DE VENTAS Y CONTROL DE INVENTARIO")
run_title.font.name = "Arial"
run_title.font.size = Pt(18)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(15, 23, 42)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = p_sub.add_run("Documentación Técnica y Proyecto Final de Software\nFecha de Entrega: 25 de septiembre de 2026")
run_sub.font.name = "Arial"
run_sub.font.size = Pt(12)
run_sub.font.italic = True

doc.add_paragraph("\n")

# Secciones del documento
def agregar_seccion(doc, titulo, contenido):
    h = doc.add_heading(titulo, level=1)
    h.style.font.name = "Arial"
    h.style.font.color.rgb = RGBColor(30, 58, 138)
    p = doc.add_paragraph(contenido)
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)

agregar_seccion(
    doc,
    "1. Introducción y Justificación",
    "El presente proyecto abarca el desarrollo e implementación de una solución de software tipo ERP (Enterprise Resource Planning) para la administración centralizada de operaciones comerciales. La aplicación integra el control de inventarios, gestión de clientes y proveedores, flujo de ventas, cuentas por cobrar/pagar y registros de auditoría mediante una arquitectura ligera y persistente."
)

agregar_seccion(
    doc,
    "2. Arquitectura del Sistema y Diseño Modular",
    "El desarrollo se fundamenta en el lenguaje Python utilizando la librería gráfica Tkinter para la capa de presentación y el motor SQLite para la capa de persistencia de datos.\n\n"
    "Estructura modular de archivos:\n"
    "• main.py: Punto de entrada principal, orquestación de la ventana contenedora y gestión de rutas entre vistas.\n"
    "• database.py: Definición de esquemas SQL, relaciones de tablas, controladores de conexión e inserción de datos iniciales.\n"
    "• config.py: Centralización de constantes del sistema, paleta de colores UI y parámetros globales.\n"
    "• Módulos específicos (ventas.py, inventario.py, clientes.py, proveedores.py, cuentas_cobrar.py, cuentas_pagar.py, reportes.py): Controladores dedicados a cada flujo del ERP."
)

agregar_seccion(
    doc,
    "3. Modelo Relacional de Base de Datos (SQL)",
    "La base de datos relacional (sistema_erp.db) está optimizada para garantizar la integridad de los datos. Contiene 6 tablas interconectadas mediante claves primarias y foráneas:\n\n"
    "1. clientes: Almacena id, nombre, tipo_documento, documento, sexo, email, telefono, direccion.\n"
    "2. proveedores: Almacena id, nombre, contacto, telefono, email.\n"
    "3. productos: Almacena id, codigo (único), nombre, precio, stock.\n"
    "4. cuentas_por_cobrar: Vinculada con la tabla clientes a través de cliente_id.\n"
    "5. cuentas_por_pagar: Vinculada con la tabla proveedores a través de proveedor_id.\n"
    "6. auditoria: Registra eventos del sistema (usuario, fecha, accion, detalles)."
)

agregar_seccion(
    doc,
    "4. Seguridad y Auditoría de Operaciones",
    "Para garantizar la trazabilidad de la información, el sistema implementa un módulo de auditoría en tiempo real. Cada transacción sensible (altas de clientes, movimientos de caja o modificaciones de stock) genera automáticamente un registro persistente con marca de tiempo (TIMESTAMP) e identificación de usuario activo."
)

agregar_seccion(
    doc,
    "5. Conclusiones y Propuesta de Escalabilidad",
    "El software cumple con los requerimientos operativos planteados. Como línea de trabajo futuro, se contempla la migración del motor SQLite hacia MySQL o PostgreSQL para entornos cliente-servidor distribuidos, así como la incorporación de gráficos analíticos dinámicos mediante Matplotlib."
)

# Guardar documento
nombre_archivo = "Trabajo_Escrito_Proyecto_ERP.docx"
doc.save(nombre_archivo)
print(f"Documento generado exitosamente: {nombre_archivo}")